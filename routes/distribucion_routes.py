from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from decimal import Decimal

from flask import render_template, request, send_file, jsonify

from sqlalchemy.orm import joinedload

from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from flask import Blueprint, send_file
from flask import redirect
from flask import url_for
from flask import flash

from flask_login import current_user, login_required

from sqlalchemy import func

from database.connection import SessionLocal

from models import prestamo
from models.periodo import Periodo
from models.configuracion import Configuracion
from models.movimiento import Movimiento
from models.socio import Socio
from models.solicitud_prestamo import SolicitudPrestamo
from models.prestamo import Prestamo
from models.transferencia import Transferencia
from services.capacidad_prestamo_service import CapacidadPrestamoService
from services.interes_service import InteresService

distribucion_bp = Blueprint(
    "distribucion",
    __name__,
    url_prefix="/distribucion"
)

@distribucion_bp.route("/<int:periodo_id>")
@login_required
def index(periodo_id):
    db = SessionLocal()
    try:
        periodo = (
            db.query(Periodo)
            .get(periodo_id)
        )

        if not periodo:

            flash(
                "Periodo no encontrado.",
                "danger"
            )

            return redirect(
                url_for("periodos.index")
            )

        config = (
            db.query(Configuracion)
            .filter(Configuracion.estado == True)
            .first()
        )

        fondo = (
            db.query(
                func.coalesce(
                    func.sum(
                        Movimiento.aporte
                    ),
                    0
                )
            )
            .filter(
                Movimiento.periodo_id == periodo.id
            )
            .scalar()
        )

        capacidad = CapacidadPrestamoService.calcular(
            db,
            periodo.id
        )["capacidad"]
        solicitudes = (

            db.query(
                SolicitudPrestamo
            )
            .options(

                joinedload(
                    SolicitudPrestamo.socio
                ),
                joinedload(
                    SolicitudPrestamo.accion
                )
            )

            .filter(

                SolicitudPrestamo.periodo_id == periodo.id,
                SolicitudPrestamo.estado == "PENDIENTE"
            )
            .order_by(

                SolicitudPrestamo.score.desc(),
                SolicitudPrestamo.prioridad.desc()
            )
            .all()
        )

        total_solicitado = sum(
            s.monto_solicitado
            for s in solicitudes

        )

        return render_template(

            "distribucion/index.html",
            periodo=periodo,
            config=config,
            fondo=fondo,
            capacidad=capacidad,
            solicitudes=solicitudes,
            total_solicitado=total_solicitado

        )

    finally:

        db.close()

@distribucion_bp.route(
    "/generar/<int:periodo_id>",
    methods=["POST"]
)
@login_required
def generar(periodo_id):

    db = SessionLocal()

    try:

        data = request.json
        print(data)
        config = (
            db.query(Configuracion)
            .filter(Configuracion.estado == True)
            .first()
        )

        if not config:
            raise Exception(
                "No existe la configuración del sistema."
            )
        
        # Validar capacidad disponible:
        total = sum(
                float(f["asignado"])
                for f in data["solicitudes"]
            )

        capacidad = CapacidadPrestamoService.calcular(
                db,
                periodo_id
            )["capacidad"]

        if total > capacidad:
                raise Exception(
                    "La distribución supera la capacidad disponible."
                )
        
        periodo = (
            db.query(Periodo)
            .filter(Periodo.id == periodo_id)
            .first()
            )

        if not periodo:
            raise Exception("El período no existe.")

        if periodo.cerrado:
            raise Exception("No se pueden generar préstamos en un período cerrado.")


        cantidad = 0
        monto = 0
        monto_total = 0

        if not data or "solicitudes" not in data:
            raise Exception("No se recibieron solicitudes.")

        for fila in data["solicitudes"]:

            print("Fila:", fila)
            monto = float(fila["asignado"])
            print("Monto:", monto)
        
            if monto <= 0:
                print("Se omite porque monto <= 0")
                continue

            solicitud = (
                db.query(SolicitudPrestamo)
                .filter(
                    SolicitudPrestamo.id == fila["id"]
                )
                .first()
            )

            if not solicitud:
                continue

            if solicitud.estado != "PENDIENTE":
                continue

            print("Solicitud:", solicitud.id)

            if monto > float(solicitud.monto_solicitado):
                raise Exception(
                    f"El monto asignado supera el solicitado para la solicitud {solicitud.id}"
                )
            
            # se omite porque ya no se puede generar préstamo si la acción ya tiene un préstamo activo    
            #existe = (
            #    db.query(Prestamo)
            #    .filter(
            #        Prestamo.accion_id == solicitud.accion_id,
            #        Prestamo.estado == "ACTIVO"
            #    )
            #    .first()
            #)

            #if existe:
            #    raise Exception(
            #        f"La acción {solicitud.accion.numero_accion} ya posee un préstamo activo."
            #    )
            
            prestamo = Prestamo(

                socio_id=solicitud.socio_id,
                accion_id=solicitud.accion_id,
                periodo_id=periodo.id,
                monto=monto,
                saldo_actual=monto,
                saldo_interes=monto,
                tasa_interes=config.interes_mensual,
                cuota_minima=0,
                fecha_prestamo=datetime.now(),                
                estado="ACTIVO"

            )
            print("Préstamo agregado")
            db.add(prestamo)

            db.flush()

            solicitud.estado = "ATENDIDA"
            solicitud.monto_aprobado = monto

            cantidad += 1
            monto_total += monto

        db.commit()

        return jsonify({

            "ok": True,
            "prestamos_generados": cantidad,
            "monto_total": monto_total,
            "mensaje": f"Se generaron {cantidad} préstamos."

        })

    except Exception as e:

        db.rollback()

        return jsonify({

            "ok":False,
            "mensaje":str(e)

        })

    finally:

        db.close()

# Ruta AJAX simular calculación de distribución de préstamos
@distribucion_bp.route(
    "/simular_distribucion",
    methods=["POST"]
)
@login_required
def simular_distribucion():

    db = SessionLocal()

    try:

        config = (
            db.query(Configuracion)
            .filter(Configuracion.estado == True)
            .first()
        )

        periodo = (
                    db.query(Periodo)
                    .filter(
                        Periodo.cerrado == False
                    )
                    .first()
                )
        
        if not periodo:
    
            flash(
                "No existe un período abierto.",
                "warning"
                    )
            return redirect(
                url_for("periodos.index")
            )
        
        socios_activos = (
            db.query(Socio)
            .filter(
                Socio.estado == True
            )
            .count()
        )

        aporte = float(config.aporte_minimo)

        ingreso_aportes = socios_activos * aporte

        amortizacion = (
            db.query(
                func.coalesce(
                    func.sum(
                        Prestamo.cuota_minima
                    ),
                    0
                )
            )
            .filter(
                Prestamo.estado=="ACTIVO"
            )
            .scalar()
        )

        prestamos = (
            db.query(Prestamo)
            .filter(
                Prestamo.estado == "ACTIVO"
            )
            .all()
        )

        intereses = sum(
            InteresService.calcular_interes(
                db,
                p.saldo_interes
            )
            for p in prestamos
        )

        #capacidad = ingreso_aportes + float(amortizacion)+float(intereses)

        capacidad=CapacidadPrestamoService.calcular(db,periodo.id)["capacidad"]
       
        solicitudes = (
            db.query(SolicitudPrestamo)
            .filter(
                SolicitudPrestamo.estado=="PENDIENTE"
            )
            .order_by(

                SolicitudPrestamo.id.asc()
            )
            .all()
        )

        restante = capacidad

        resultado = []

        for s in solicitudes:

            solicitado = float(
                s.monto_solicitado
            )

            aprobado = min(
                solicitado,
                restante
            )

            resultado.append({

                "socio": s.socio.nombres,
                "solicitado": solicitado,
                "aprobado": aprobado,
                "restante": restante-aprobado

            })

            restante -= aprobado

            if restante <= 0:

                break

        return render_template(

            "distribucion/index.html",
            capacidad=capacidad,
            resultado=resultado,
            restante=restante

        )

    finally:

        db.close()


@distribucion_bp.route(
    "/revertir/<int:periodo_id>",
    methods=["POST"]
)
@login_required
def revertir(periodo_id):

    db = SessionLocal()

    try:

        periodo = (
            db.query(Periodo)
            .filter(
                Periodo.id == periodo_id
            )
            .first()
        )

        if not periodo:
            raise Exception(
                "El período no existe."
            )

        prestamos = (
            db.query(Prestamo)
            .filter(
                Prestamo.periodo_id == periodo_id,
                Prestamo.estado == "ACTIVO"
            )
            .all()
        )

        if not prestamos:
            raise Exception(
                "No existen préstamos generados para este período."
            )

        cantidad = 0
        monto_total = 0

        for prestamo in prestamos:

            #
            # Verificar que no tenga pagos
            #

            movimientos = (
                db.query(Movimiento)
                .filter(
                    Movimiento.prestamo_id == prestamo.id
                )
                .count()
            )

            if movimientos > 0:

                raise Exception(
                    f"El préstamo {prestamo.id} ya tiene movimientos registrados."
                )
            
            solicitud = (
                db.query(SolicitudPrestamo)
                .filter(
                    SolicitudPrestamo.periodo_id == periodo.id,
                    SolicitudPrestamo.socio_id == prestamo.socio_id,
                    SolicitudPrestamo.accion_id == prestamo.accion_id,
                    SolicitudPrestamo.estado == "ATENDIDA"
                )
                .first()
            )

            if solicitud:

                solicitud.estado = "PENDIENTE"
                solicitud.monto_aprobado = 0
                #solicitud.prestamo_id = None

            monto_total += float(prestamo.monto)
            prestamo.estado = "ANULADO"
            cantidad += 1

        db.commit()

        return jsonify({

            "ok": True,
            "cantidad": cantidad,
            "monto_total": monto_total,
            "mensaje": f"Se revirtieron {cantidad} préstamos."

        })

    except Exception as e:

        db.rollback()

        return jsonify({

            "ok": False,
            "mensaje": str(e)

        })

    finally:

        db.close()


@distribucion_bp.route(
    "/simulacion_transferencias/<int:periodo_id>"
)
@login_required
def simulacion_transferencias(periodo_id):

    db = SessionLocal()

    try:

        # =====================================================
        # PERÍODO
        # =====================================================

        periodo = (
            db.query(Periodo)
            .filter(Periodo.id == periodo_id)
            .first()
        )

        if not periodo:
            flash("Período no encontrado.", "danger")
            return redirect(
                url_for("distribucion.index")
            )

        # =====================================================
        # SOLICITUDES APROBADAS
        # =====================================================

        solicitudes = (
            db.query(SolicitudPrestamo)
            .options(
                joinedload(SolicitudPrestamo.socio),
                joinedload(SolicitudPrestamo.accion)
            )
            .filter(
                SolicitudPrestamo.periodo_id == periodo_id,
                SolicitudPrestamo.estado == "ATENDIDA",
                SolicitudPrestamo.monto_aprobado > 0
            )
            .all()
        )

        # =====================================================
        # AGRUPAR PRÉSTAMOS POR SOCIO
        # =====================================================

        socios_prestamos = defaultdict(
            lambda: {
                "socio_id": None,
                "nombre": "",
                "prestamos": [],
                "total_aprobado": Decimal("0.00")
            }
        )

        for s in solicitudes:

            socio_id = s.socio_id
            item = socios_prestamos[socio_id]
            item["socio_id"] = socio_id
            item["nombre"] = (
                s.socio.nombres
                if s.socio
                else f"Socio #{socio_id}"
            )

            monto = Decimal(
                str(s.monto_aprobado or 0)
            )

            item["prestamos"].append({
                "solicitud_id": s.id,
                "accion_id": s.accion_id,
                "accion": (
                    s.accion.numero_accion
                    if s.accion
                    else ""
                ),
                "monto": float(monto)
            })

            item["total_aprobado"] += monto

        # =====================================================
        # FONDOS DEL PERÍODO
        # =====================================================

        movimientos = (
            db.query(Movimiento)
            .filter(
                Movimiento.periodo_id == periodo_id
            )
            .all()
        )

        fondos = defaultdict(
            lambda: Decimal("0.00")
        )

        nombres_socios = {}

        for m in movimientos:

            socio_id = m.socio_id

            fondos[socio_id] += Decimal(str(m.cuota_pagada or 0)) + Decimal(str(m.sobre or 0))

            if m.socio:

                nombres_socios[socio_id] = (
                    m.socio.nombres
                )

        # =====================================================
        # UNIR SOCIOS CON PRÉSTAMOS + SOCIOS CON FONDOS
        # =====================================================

        todos_socios_ids = set(
            socios_prestamos.keys()
        ) | set(
            fondos.keys()
        )

        socios = []

        for socio_id in todos_socios_ids:

            prestamo_data = socios_prestamos.get(
                socio_id,
                {
                    "nombre": "",
                    "prestamos": [],
                    "total_aprobado": Decimal("0.00")
                }
            )

            nombre = (
                prestamo_data["nombre"]
                or nombres_socios.get(
                    socio_id,
                    f"Socio #{socio_id}"
                )
            )

            total_aprobado = Decimal(
                str(
                    prestamo_data["total_aprobado"]
                )
            )

            fondo_propio = fondos.get(
                socio_id,
                Decimal("0.00")
            )

            # =================================================
            # APLICAR FONDO PROPIO
            # =================================================

            propio_aplicado = min(
                fondo_propio,
                total_aprobado
            )

            # =================================================
            # NECESIDAD
            # =================================================

            necesidad = max(
                Decimal("0.00"),
                total_aprobado
                - propio_aplicado
            )

            # =================================================
            # EXCEDENTE
            # =================================================

            excedente = max(
                Decimal("0.00"),
                fondo_propio
                - propio_aplicado
            )

            socios.append({
                "socio_id": socio_id,
                "nombre": nombre,
                "prestamos":prestamo_data["prestamos"],
                "total_aprobado":float(total_aprobado),
                "cuota_propia":float(fondo_propio),
                "propio_aplicado":float(propio_aplicado),
                "necesidad":float(necesidad),
                "excedente":float(excedente)
            })

        # =====================================================
        # ORDENAR POR NOMBRE
        # =====================================================

        socios.sort(
            key=lambda x:
                x["nombre"].upper()
        )

        # =====================================================
        # FONDO TOTAL
        # =====================================================
        fondo = sum(
            (
                fondos.get(
                    socio_id,
                    Decimal("0.00")
                )
                for socio_id in fondos
            ),
            Decimal("0.00")
        )

        # =====================================================
        # TOTAL APROBADO
        # =====================================================

        total_aprobado = sum(
            (
                Decimal(
                    str(
                        s["total_aprobado"]
                    )
                )
                for s in socios
            ),
            Decimal("0.00")
        )

        # =====================================================
        # ENVIAR A TEMPLATE
        # =====================================================

        return render_template(
            "distribucion/simulacion_transferencias.html",
            periodo_id=periodo_id,
            periodo=periodo,
            solicitudes=solicitudes,
            socios=socios,
            fondo=fondo,
            total_aprobado=total_aprobado
        )

    finally:

        db.close()

@distribucion_bp.route(
    "/transferencias/confirmar/<int:periodo_id>",
    methods=["POST"]
)
@login_required
def confirmar_transferencias(periodo_id):

    db = SessionLocal()

    try:

        data = request.get_json() or {}
        transferencias = data.get("transferencias", [])

        if not transferencias:

            return jsonify({
                "ok": False,
                "mensaje":
                    "No existen transferencias para confirmar."
            }), 400


        periodo = (
            db.query(Periodo)
            .filter(
                Periodo.id == periodo_id
            )
            .first()
        )

        if not periodo:

            return jsonify({
                "ok": False,
                "mensaje":
                    "Período no encontrado."
            }), 404

        if periodo.cerrado:

            return jsonify({
                "ok": False,
                "mensaje":
                    "El período ya está cerrado."
            }), 400


        # =====================================================
        # EVITAR DUPLICAR TRANSFERENCIAS
        # =====================================================

        existentes = (
            db.query(Transferencia)
            .filter(
                Transferencia.periodo_id ==
                periodo_id,

                Transferencia.estado ==
                "CONFIRMADA"
            )
            .count()
        )


        if existentes > 0:

            return jsonify({
                "ok": False,
                "mensaje":
                    "El período ya tiene transferencias confirmadas."
            }), 400


        total = Decimal("0.00")


        for item in transferencias:

            origen_id = int(
                item["socio_origen_id"]
            )

            destino_id = int(
                item["socio_destino_id"]
            )

            monto = Decimal(
                str(item["monto"])
            )


            if origen_id == destino_id:

                raise Exception(
                    "No se permite transferir "
                    "a sí mismo."
                )


            if monto <= 0:

                raise Exception(
                    "El monto de una transferencia "
                    "debe ser mayor a cero."
                )

            transferencia = Transferencia(
                periodo_id=periodo_id,
                socio_origen_id= origen_id,
                socio_destino_id= destino_id,
                monto=monto,
                estado="CONFIRMADA",
                usuario_id=current_user.id
            )

            db.add(
                transferencia
            )

            total += monto

        db.commit()

        return jsonify({

            "ok": True,

            "mensaje":
                (
                    f"Se confirmaron "
                    f"{len(transferencias)} "
                    f"transferencias por "
                    f"S/ {total:,.2f}."
                ),

            "cantidad":
                len(transferencias),

            "total":
                float(total),
            "redirect_url": f"/distribucion/simulacion_transferencias/{periodo_id}"  # Corregido con F-string
        })


    except Exception as e:

        db.rollback()

        return jsonify({
            "ok": False,
            "mensaje": str(e)

        }), 400

    finally:

        db.close()
@distribucion_bp.route(
    "/transferencias/visor/<int:periodo_id>"
)
@login_required
def visor_transferencias(periodo_id):

    db = SessionLocal()

    try:

        # =====================================================
        # PERÍODO
        # =====================================================

        periodo = (
            db.query(Periodo)
            .filter(
                Periodo.id == periodo_id
            )
            .first()
        )

        if not periodo:
            return "Período no encontrado", 404


        # =====================================================
        # TRANSFERENCIAS CONFIRMADAS
        # =====================================================

        transferencias = (
            db.query(Transferencia)
            .options(
                joinedload(Transferencia.socio_origen),
                joinedload(Transferencia.socio_destino)
            )
            .filter(
                Transferencia.periodo_id == periodo_id,
                Transferencia.estado == "CONFIRMADA"
            )
            .order_by(
                Transferencia.socio_destino_id,
                Transferencia.socio_origen_id
            )
            .all()
        )


        # =====================================================
        # SOLICITUDES ATENDIDAS
        #
        # Se utilizan solamente para determinar el aporte
        # propio aplicado. NO se muestra el monto aprobado
        # como monto a recibir en el visor.
        # =====================================================

        solicitudes = (
            db.query(SolicitudPrestamo)
            .options(
                joinedload(
                    SolicitudPrestamo.socio
                )
            )
            .filter(
                SolicitudPrestamo.periodo_id == periodo_id,
                SolicitudPrestamo.estado == "ATENDIDA",
                SolicitudPrestamo.monto_aprobado > 0
            )
            .all()
        )

        # =====================================================
        # APORTE APROBADO POR SOCIO
        # Se conserva internamente para no aplicar como
        # "propio" un importe superior al préstamo aprobado.
        # =====================================================
        aprobados = defaultdict(
            lambda: Decimal("0.00")
        )

        nombres = {}

        for solicitud in solicitudes:

            socio_id = solicitud.socio_id

            monto_aprobado = Decimal(
                str(
                    solicitud.monto_aprobado or 0
                )
            )

            aprobados[socio_id] += (
                monto_aprobado
            )

            if solicitud.socio:

                nombres[socio_id] = (
                    solicitud.socio.nombres
                )

            else:

                nombres[socio_id] = (
                    f"Socio #{socio_id}"
                )


        # =====================================================
        # MOVIMIENTOS DEL PERÍODO
        #
        # Aporte propio:
        # cuota_pagada + sobre
        # =====================================================

        movimientos = (
            db.query(Movimiento)
            .filter(
                Movimiento.periodo_id == periodo_id
            )
            .all()
        )

        fondos_propios = defaultdict(
            lambda: Decimal("0.00")
        )

        for movimiento in movimientos:

            cuota = Decimal(str(movimiento.cuota_pagada or 0))
            sobre = Decimal(str(movimiento.sobre or 0))
            fondos_propios[
                movimiento.socio_id
            ] += cuota + sobre


        # =====================================================
        # APORTE PROPIO APLICADO
        # Se limita al monto aprobado para evitar que el
        # aporte propio genere un total financiado superior
        # al préstamo correspondiente.
        # =====================================================

        propio_aplicado = {}


        for socio_id, aprobado in aprobados.items():

            propio = fondos_propios.get(
                socio_id,
                Decimal("0.00")
            )

            propio_aplicado[socio_id] = min(
                propio,
                aprobado
            )

        # =====================================================
        # AGRUPAR TRANSFERENCIAS POR BENEFICIARIO
        # =====================================================

        transferencias_por_beneficiario = (
            defaultdict(list)
        )

        for transferencia in transferencias:

            transferencias_por_beneficiario[
                transferencia.socio_destino_id
            ].append(
                transferencia
            )


        # =====================================================
        # SOCIOS QUE APARECERÁN EN EL VISOR
        # =====================================================

        socio_ids = set(
            aprobados.keys()
        )


        for transferencia in transferencias:

            socio_ids.add(
                transferencia.socio_destino_id
            )


        # =====================================================
        # CONSTRUIR VISOR
        # =====================================================

        beneficiarios = []


        for socio_id in sorted(socio_ids):

            nombre_beneficiario = nombres.get(
                socio_id,
                f"Socio #{socio_id}"
            )

            lista_transferencias = (
                transferencias_por_beneficiario.get(
                    socio_id,
                    []
                )
            )


            # =================================================
            # APORTE PROPIO
            # =================================================

            aporte_propio = propio_aplicado.get(
                socio_id,
                Decimal("0.00")
            )

            # =================================================
            # TRANSFERENCIAS RECIBIDAS
            # =================================================

            total_transferencias = sum(
                (
                    Decimal(str(t.monto or 0))
                    for t in lista_transferencias
                ),
                Decimal("0.00")
            )


            # =================================================
            # TOTAL FINAL
            #
            # APORTE PROPIO + TRANSFERENCIAS
            # =================================================
            total_final = (
                aporte_propio +
                total_transferencias
            )

            # =================================================
            # DETALLE
            # =================================================

            filas = []

            # -------------------------------------------------
            # APORTE PROPIO
            # -------------------------------------------------

            if aporte_propio > 0:

                filas.append({

                    "origen": nombre_beneficiario,
                    "monto": float(aporte_propio),
                    "tipo": "APORTE PROPIO"

                })


            # -------------------------------------------------
            # TRANSFERENCIAS
            # -------------------------------------------------
            for transferencia in lista_transferencias:

                origen_id = (
                    transferencia.socio_origen_id
                )


                origen_nombre = (

                    transferencia.socio_origen.nombres

                    if transferencia.socio_origen

                    else nombres.get(
                        origen_id,
                        f"Socio #{origen_id}"
                    )

                )

                filas.append({

                    "origen": origen_nombre,
                    "monto": float(transferencia.monto or 0),
                    "tipo": "TRANSFERENCIA"

                })


            # =================================================
            # AGREGAR BENEFICIARIO
            # =================================================
            beneficiarios.append({

                "socio_id": socio_id,
                "beneficiario": nombre_beneficiario,
                "aporte_propio": float(aporte_propio),
                "transferencias": float(total_transferencias),
                "total_final":float(total_final),
                "filas":filas

            })


        # =====================================================
        # TOTALES GENERALES
        # =====================================================

        total_aporte_propio = sum(
            (
                Decimal(str(x["aporte_propio"])
                )
                for x in beneficiarios
            ),
            Decimal("0.00")
        )


        total_transferencias = sum(
            (
                Decimal(str(x["transferencias"])
                )
                for x in beneficiarios
            ),
            Decimal("0.00")
        )


        total_final = (
            total_aporte_propio +
            total_transferencias
        )


        # =====================================================
        # RENDER
        # =====================================================

        return render_template(
            "distribucion/visor_transferencias.html",

            periodo=periodo,
            beneficiarios=beneficiarios,
            total_aporte_propio=float(total_aporte_propio),
            total_transferencias=float(total_transferencias),
            total_final=float(total_final),
            cantidad_transferencias=len(transferencias)
        )

    finally:

        db.close()

    
@distribucion_bp.route(
    "/transferencias/exportar/<int:periodo_id>"
)
@login_required
def exportar_transferencias_excel(periodo_id):
    
    db = SessionLocal()

    try:

        # =====================================================
        # PERÍODO
        # =====================================================

        periodo = (
            db.query(Periodo)
            .filter(
                Periodo.id == periodo_id
            )
            .first()
        )

        if not periodo:

            return "Período no encontrado", 404


        # =====================================================
        # TRANSFERENCIAS CONFIRMADAS
        # =====================================================

        transferencias = (
            db.query(Transferencia)
            .options(
                joinedload(
                    Transferencia.socio_origen
                ),
                joinedload(
                    Transferencia.socio_destino
                )
            )
            .filter(
                Transferencia.periodo_id == periodo_id,
                Transferencia.estado == "CONFIRMADA"
            )
            .order_by(
                Transferencia.socio_destino_id,
                Transferencia.socio_origen_id
            )
            .all()
        )


        # =====================================================
        # SOLICITUDES DE PRÉSTAMO APROBADAS
        #
        # Estas sirven únicamente para conocer:
        # cuánto fue aprobado para cada socio.
        #
        # NO se exporta monto_solicitado.
        # =====================================================

        solicitudes = (
            db.query(SolicitudPrestamo)
            .options(
                joinedload(
                    SolicitudPrestamo.socio
                )
            )
            .filter(
                SolicitudPrestamo.periodo_id == periodo_id,
                SolicitudPrestamo.estado == "ATENDIDA",
                SolicitudPrestamo.monto_aprobado > 0
            )
            .all()
        )


        aprobados = defaultdict(
            lambda: Decimal("0.00")
        )

        nombres = {}


        for s in solicitudes:

            socio_id = s.socio_id

            monto_aprobado = Decimal(
                str(
                    s.monto_aprobado
                    or 0
                )
            )

            aprobados[socio_id] += (
                monto_aprobado
            )

            if s.socio:

                nombres[socio_id] = (
                    s.socio.nombres
                )

            else:

                nombres[socio_id] = (
                    f"Socio #{socio_id}"
                )


        # =====================================================
        # APORTE PROPIO
        #
        # Se calcula:
        #
        # cuota_pagada + sobre
        #
        # y se limita al monto aprobado.
        #
        # IMPORTANTE:
        # Este dinero NO es una transferencia.
        # Es solamente referencial para el tesorero.
        # =====================================================

        movimientos = (
            db.query(Movimiento)
            .filter(
                Movimiento.periodo_id == periodo_id
            )
            .all()
        )

        fondos = defaultdict(
            lambda: Decimal("0.00")
        )


        for m in movimientos:

            cuota = Decimal(
                str(
                    m.cuota_pagada
                    or 0
                )
            )

            sobre = Decimal(
                str(
                    m.sobre
                    or 0
                )
            )

            fondos[m.socio_id] += (
                cuota + sobre
            )


        aporte_propio = {}


        for socio_id, monto_aprobado in aprobados.items():

            aporte = fondos.get(
                socio_id,
                Decimal("0.00")
            )

            aporte_propio[socio_id] = min(
                aporte,
                monto_aprobado
            )


        # =====================================================
        # AGRUPAR TRANSFERENCIAS POR BENEFICIARIO
        # =====================================================

        grupos = defaultdict(list)

        for t in transferencias:

            grupos[
                t.socio_destino_id
            ].append(t)


        # =====================================================
        # ASEGURAR QUE TAMBIÉN APAREZCAN SOCIOS
        # QUE TIENEN APORTE PROPIO
        #
        # Aunque por algún motivo tengan transferencia 0,
        # deben aparecer en el reporte.
        # =====================================================

        beneficiarios_ids = set(
            grupos.keys()
        )

        beneficiarios_ids.update(
            socio_id
            for socio_id, monto in aporte_propio.items()
            if monto > 0
        )


        # =====================================================
        # CREAR EXCEL
        # =====================================================

        wb = Workbook()
        ws = wb.active
        ws.title = "Transferencias"


        # =====================================================
        # CABECERA INFORMATIVA
        # =====================================================

        ws.merge_cells(
            "A1:F1"
        )

        ws["A1"] = (
            "REPORTE DE TRANSFERENCIAS "
            "Y APORTE PROPIO"
        )

        ws["A1"].font = Font(
            bold=True,
            size=14
        )

        ws["A1"].alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


        ws.merge_cells(
            "A2:F2"
        )

        ws["A2"] = (
            f"Período: "
            f"{periodo.anio}-"
            f"{periodo.mes:02d}"
        )

        ws["A2"].font = Font(
            italic=True
        )

        ws["A2"].alignment = Alignment(
            horizontal="center"
        )

        # =====================================================
        # CABECERA DE TABLA
        # =====================================================

        fila_inicio = 4

        encabezados = [
            "Beneficiario",
            "Origen / Detalle",
            "Moneda",
            "Monto",
            "Tipo de Fila",
            "Observación"
        ]


        for columna, valor in enumerate(
            encabezados,
            start=1
        ):

            cell = ws.cell(
                row=fila_inicio,
                column=columna,
                value=valor
            )

            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="198754"
            )


        fila = fila_inicio + 1

        total_transferencias = Decimal(
            "0.00"
        )

        total_aporte_propio = Decimal(
            "0.00"
        )

        total_general = Decimal(
            "0.00"
        )


        # =====================================================
        # GENERAR REPORTE
        # =====================================================

        for socio_id in sorted(
            beneficiarios_ids,
            key=lambda x: nombres.get(
                x,
                f"Socio #{x}"
            ).lower()
        ):

            lista_transferencias = grupos.get(
                socio_id,
                []
            )


            beneficiario = nombres.get(
                socio_id
            )


            if not beneficiario:

                if lista_transferencias:

                    beneficiario = (
                        lista_transferencias[0]
                        .socio_destino.nombres
                        if lista_transferencias[0]
                        .socio_destino
                        else f"Socio #{socio_id}"
                    )

                else:

                    beneficiario = (
                        f"Socio #{socio_id}"
                    )


            # =================================================
            # TOTAL DE TRANSFERENCIAS DEL BENEFICIARIO
            # =================================================

            subtotal_transferencias = Decimal(
                "0.00"
            )


            # =================================================
            # APORTE PROPIO
            # =================================================

            monto_propio = aporte_propio.get(
                socio_id,
                Decimal("0.00")
            )


            # =================================================
            # FILAS DE TRANSFERENCIAS
            # =================================================

            for indice, t in enumerate(
                lista_transferencias
            ):

                origen = (
                    t.socio_origen.nombres
                    if t.socio_origen
                    else
                    f"Socio #{t.socio_origen_id}"
                )


                monto = Decimal(
                    str(
                        t.monto
                        or 0
                    )
                )

                # Para que el beneficiario
                # no se repita visualmente.
                beneficiario_excel = (
                    beneficiario
                    if indice == 0
                    else ""
                )


                ws.append([
                    beneficiario_excel,
                    origen,
                    "S/",
                    float(monto),
                    "TRANSFERENCIA",
                    "Monto que debe transferirse"
                ])


                subtotal_transferencias += (
                    monto
                )

                total_transferencias += (
                    monto
                )

                fila += 1


            # =================================================
            # APORTE PROPIO REFERENCIAL
            # =================================================

            if monto_propio > 0:

                # Si no hubo transferencia,
                # ponemos el nombre.
                beneficiario_excel = (
                    beneficiario
                    if not lista_transferencias
                    else ""
                )


                ws.append([
                    beneficiario_excel,
                    "APORTE PROPIO DEL SOCIO",
                    "S/",
                    float(monto_propio),
                    "APORTE PROPIO",
                    "Referencial - no transferir"
                ])


                total_aporte_propio += (
                    monto_propio
                )

                fila += 1


            # =================================================
            # TOTAL A RECIBIR
            #
            # Transferencias + aporte propio
            # =================================================

            total_beneficiario = (
                subtotal_transferencias
                +
                monto_propio
            )


            total_general += (
                total_beneficiario
            )


            ws.append([
                "",
                "TOTAL A RECIBIR",
                "S/",
                float(total_beneficiario),
                "TOTAL",
                "Transferencias + aporte propio"
            ])


            # =================================================
            # ESTILO DEL TOTAL
            # =================================================

            for columna in range(
                1,
                7
            ):

                cell = ws.cell(
                    row=fila,
                    column=columna
                )

                cell.font = Font(
                    bold=True
                )

                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor="D1E7DD"
                )


            fila += 1


        # =====================================================
        # RESUMEN FINAL
        # =====================================================

        fila += 1


        ws.append([
            "RESUMEN GENERAL",
            "",
            "",
            "",
            "",
            ""
        ])


        fila_resumen = fila


        for columna in range(
            1,
            7
        ):

            cell = ws.cell(
                row=fila_resumen,
                column=columna
            )

            cell.font = Font(
                bold=True
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="198754"
            )

            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )


        fila += 1


        ws.append([
            "",
            "TOTAL TRANSFERENCIAS",
            "S/",
            float(total_transferencias),
            "TRANSFERENCIAS",
            "Dinero que debe transferirse"
        ])


        fila += 1


        ws.append([
            "",
            "TOTAL APORTE PROPIO",
            "S/",
            float(total_aporte_propio),
            "APORTE PROPIO",
            "Dinero propio del socio"
        ])


        fila += 1


        ws.append([
            "",
            "TOTAL GENERAL A RECIBIR",
            "S/",
            float(total_general),
            "TOTAL",
            "Transferencias + aporte propio"
        ])


        # =====================================================
        # FORMATO NUMÉRICO
        # =====================================================

        for row in range(
            fila_inicio + 1,
            ws.max_row + 1
        ):

            ws.cell(
                row=row,
                column=4
            ).number_format = (
                '#,##0.00'
            )


        # =====================================================
        # ALINEACIÓN
        # =====================================================

        for row in ws.iter_rows():

            for cell in row:

                cell.alignment = Alignment(
                    vertical="center"
                )


        # =====================================================
        # ALINEACIÓN DE MONTOS
        # =====================================================

        for row in range(
            fila_inicio + 1,
            ws.max_row + 1
        ):

            ws.cell(
                row=row,
                column=3
            ).alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            ws.cell(
                row=row,
                column=4
            ).alignment = Alignment(
                horizontal="right",
                vertical="center"
            )


        # =====================================================
        # ANCHOS
        # =====================================================

        anchos = {
            "A": 35,
            "B": 38,
            "C": 10,
            "D": 18,
            "E": 22,
            "F": 35
        }


        for columna, ancho in anchos.items():

            ws.column_dimensions[
                columna
            ].width = ancho


        # =====================================================
        # FILTRO
        # =====================================================

        ws.auto_filter.ref = (
            f"A{fila_inicio}:F{fila - 1}"
        )


        # =====================================================
        # CONGELAR CABECERA
        # =====================================================

        ws.freeze_panes = "A5"


        # =====================================================
        # ALTURA DE FILAS
        # =====================================================

        ws.row_dimensions[1].height = 25

        ws.row_dimensions[4].height = 30


        # =====================================================
        # SEGUNDA HOJA: RESUMEN PARA TESORERÍA
        # =====================================================

        ws2 = wb.create_sheet(
            "Resumen Tesorería"
        )


        ws2.append([
            "BENEFICIARIO",
            "TRANSFERENCIAS",
            "APORTE PROPIO",
            "TOTAL A RECIBIR"
        ])


        for cell in ws2[1]:

            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="198754"
            )

            cell.alignment = Alignment(
                horizontal="center"
            )


        for socio_id in sorted(
            beneficiarios_ids,
            key=lambda x: nombres.get(
                x,
                f"Socio #{x}"
            ).lower()
        ):

            beneficiario = nombres.get(
                socio_id,
                f"Socio #{socio_id}"
            )


            total_trans = sum(
                (
                    Decimal(
                        str(
                            t.monto
                            or 0
                        )
                    )
                    for t in grupos.get(
                        socio_id,
                        []
                    )
                ),
                Decimal("0.00")
            )


            propio_socio = aporte_propio.get(
                socio_id,
                Decimal("0.00")
            )


            total_recibir = (
                total_trans
                +
                propio_socio
            )


            ws2.append([
                beneficiario,
                float(total_trans),
                float(propio_socio),
                float(total_recibir)
            ])


        # =====================================================
        # TOTALES RESUMEN
        # =====================================================

        ws2.append([
            "TOTAL GENERAL",
            float(total_transferencias),
            float(total_aporte_propio),
            float(total_general)
        ])


        ultima_fila = ws2.max_row


        for cell in ws2[ultima_fila]:

            cell.font = Font(
                bold=True
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="D1E7DD"
            )


        for row in ws2.iter_rows(
            min_row=2,
            max_row=ws2.max_row,
            min_col=2,
            max_col=4
        ):

            for cell in row:

                cell.number_format = (
                    '#,##0.00'
                )

                cell.alignment = Alignment(
                    horizontal="right"
                )


        ws2.column_dimensions["A"].width = 38
        ws2.column_dimensions["B"].width = 20
        ws2.column_dimensions["C"].width = 20
        ws2.column_dimensions["D"].width = 20


        ws2.freeze_panes = "A2"


        # =====================================================
        # GENERAR ARCHIVO
        # =====================================================

        output = BytesIO()

        wb.save(output)

        output.seek(0)


        nombre_archivo = (
            f"transferencias_"
            f"{periodo.anio}_"
            f"{periodo.mes:02d}.xlsx"
        )


        return send_file(
            output,
            as_attachment=True,
            download_name=nombre_archivo,
            mimetype=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            )
        )


    finally:

        db.close()

# ============================================================
# REPORTE ESTADO DE OPERACIONES
# ============================================================

@distribucion_bp.route("/reporte_estado_operaciones/<int:periodo_id>")
@login_required
def reporte_estado_operaciones(periodo_id):
    db = SessionLocal()
    try:
        periodo = (
            db.query(Periodo)
            .get(periodo_id)
        )

        if not periodo:

            flash(
                "Periodo no encontrado.",
                "danger"
            )

            return redirect(
                url_for("periodos.index")
            )

        config = (
            db.query(Configuracion)
            .filter(Configuracion.estado == True)
            .first()
        )

        return render_template(
            "distribucion/estado_operaciones.html",

            periodo=periodo,
            config=config
        )

    finally:

        db.close()

#
@distribucion_bp.route(
    "/estado-operaciones/exportar/<int:periodo_id>"
)
@login_required
def exportar_estado_operaciones_excel(periodo_id):

    db = SessionLocal()

    try:

        # ====================================================
        # PERÍODO
        # ====================================================

        periodo_actual = (
            db.query(Periodo)
            .filter(
                Periodo.id == periodo_id
            )
            .first()
        )

        if not periodo_actual:

            return (
                "Período no encontrado",
                404
            )

        anio = periodo_actual.anio

        # ====================================================
        # SOCIOS
        # ====================================================
        socios = (
            db.query(Socio)
            .options(
                joinedload(
                    Socio.acciones
                )
            )
            .filter(
                Socio.estado == True
            )
            .order_by(
                Socio.nombres
            )
            .all()
        )


        # ====================================================
        # PERÍODOS DEL AÑO
        #
        # Se cargan los 12 meses para que también aparezcan
        # los meses futuros aunque todavía no tengan movimientos.
        # ====================================================

        periodos = (
            db.query(Periodo)
            .filter(
                Periodo.anio == anio
            )
            .order_by(
                Periodo.mes
            )
            .all()
        )


        periodos_por_mes = {
            p.mes: p
            for p in periodos
        }


        # ====================================================
        # MOVIMIENTOS
        #
        # IMPORTANTE:
        # Se cargan todos los movimientos del año pero después
        # se agrupan por ACCION.
        #
        # Esto evita mezclar los movimientos de Acción 1,
        # Acción 2, etc. del mismo socio.
        # ====================================================

        movimientos = (
            db.query(Movimiento)
            .filter(
                Movimiento.periodo_id.in_(
                    [p.id for p in periodos]
                )
            )
            .all()
        )

        movimientos_por_accion_mes = defaultdict(
            lambda: defaultdict(list)
        )

        for movimiento in movimientos:

            if movimiento.accion_id is None:
                continue

            periodo_mov = (
                next(
                    (
                        p
                        for p in periodos
                        if p.id ==
                        movimiento.periodo_id
                    ),
                    None
                )
            )

            if not periodo_mov:
                continue

            movimientos_por_accion_mes[
                movimiento.accion_id
            ][
                periodo_mov.mes
            ].append(
                movimiento
            )

        # ====================================================
        # PRÉSTAMOS
        # MUY IMPORTANTE:
        # Se usa Prestamo.monto.
        # porque este reporte representa el estado real
        # de operaciones.
        # ====================================================
        prestamos = (
            db.query(Prestamo)
            .filter(
                Prestamo.periodo_id.in_(
                    [p.id for p in periodos]
                )
            )
            .all()
        )

        prestamos_por_accion = defaultdict(list)

        for prestamo in prestamos:

            if prestamo.accion_id is None:
                continue

            prestamos_por_accion[
                prestamo.accion_id
            ].append(
                prestamo
            )

        prestamos_por_mes = {
            mes: Decimal("0.00")
            for mes in range(1, 13)
        }

        cuota_fija_por_mes = {
            mes: Decimal("0.00")
            for mes in range(1, 13)
        }     

        for prestamo in prestamos:

            periodo_prestamo = periodos_por_mes.get(
                next(
                    (
                        p.mes
                        for p in periodos
                        if p.id == prestamo.periodo_id
                    ),
                    None
                )
            )

            if not periodo_prestamo:
                continue

            mes_prestamo = periodo_prestamo.mes

            prestamos_por_mes[mes_prestamo] += Decimal(
                prestamo.monto
            )    

        for mes in range(1, 13):

            periodo_mes = periodos_por_mes.get(mes)

            if not periodo_mes:
                continue

            for prestamo in prestamos:

                if prestamo.periodo_id > periodo_mes.id:
                    continue

                if prestamo.estado == "ANULADO":
                    continue

                cuota_fija_por_mes[mes] += Decimal(
                    prestamo.cuota_minima
                )


        saldo_por_mes = {
            mes: Decimal("0.00")
            for mes in range(1, 13)
        }

        for accion_id, movimientos_meses in movimientos_por_accion_mes.items():

            saldo_accion = Decimal("0.00")

            for mes in range(1, 13):

                lista_mes = movimientos_meses.get(
                    mes,
                    []
                )

                if lista_mes:

                    ultimo_mov = max(
                        lista_mes,
                        key=lambda x: (
                            x.fecha_registro
                            or 0
                        )
                    )

                    saldo_accion = Decimal(
                        ultimo_mov.saldo_prestamo
                    )

                saldo_por_mes[mes] += saldo_accion

        # ====================================================
        # UTILIDADES DECIMALES
        # ====================================================

        def suma_movimientos(
            lista,
            campo
        ):

            total = Decimal("0.00")

            for movimiento in lista:

                total += Decimal(
                    getattr(
                        movimiento,
                        campo,
                        0
                    )
                )

            return total


        # ====================================================
        # CREAR EXCEL
        # ====================================================

        wb = Workbook()
        ws = wb.active
        ws.title = "Estado de Operaciones"


        # ====================================================
        # COLORES
        # ====================================================

        azul_oscuro = "17365D"
        azul = "5B9BD5"
        azul_claro = "D9EAF7"
        verde = "92D050"
        verde_claro = "E2F0D9"
        celeste = "00B0F0"
        amarillo = "FFF2CC"
        gris = "D9E1F2"
        blanco = "FFFFFF"
        negro = "000000"
        rojo = "FF0000"
        thin_black = Side(style="thin",color="000000")
        border = Border(left=thin_black,right=thin_black,top=thin_black,bottom=thin_black)

        # ====================================================
        # TÍTULO GENERAL
        # ====================================================
        ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=11)

        titulo = ws.cell(row=1,column=1)

        titulo.value = (f"BANQUITO FAMILIAR " f"(ESTADO DE OPERACIONES {anio})")

        titulo.font = Font(bold=True,color=blanco,size=14)
        titulo.fill = PatternFill(fill_type="solid",fgColor=azul_oscuro)
        titulo.alignment = Alignment(horizontal="center",vertical="center")
        ws.row_dimensions[1].height = 25

        fila = 3

        # ====================================================
        # TOTALES GENERALES
        # ====================================================
        gran_total_aportes = Decimal("0.00")
        gran_total_amortizacion = Decimal("0.00")
        gran_total_intereses = Decimal("0.00")
        gran_total = Decimal("0.00")
        gran_total_prestamos = Decimal("0.00")
        gran_total_multas = Decimal("0.00")

        # ====================================================
        # RESUMEN MENSUAL GENERAL
        # ====================================================
        resumen_mensual = {
            mes: {
                "aportes": Decimal("0.00"),
                "amortizacion": Decimal("0.00"),
                "intereses": Decimal("0.00"),
                "total": Decimal("0.00"),
                "cuota_pagar": Decimal("0.00"),
                "cuota_fija": Decimal("0.00"),
                "prestamos": Decimal("0.00"),
                "multa": Decimal("0.00"),
                "sobre": Decimal("0.00"),
                "saldo_prestamos": Decimal("0.00"),
            }
            for mes in range(1, 13)
        }

        ultimo_saldo_accion_mes = defaultdict(
            lambda: Decimal("0.00")
        )



        # ====================================================
        # RECORRER SOCIOS
        # ====================================================

        for socio in socios:

            acciones = sorted(
                socio.acciones or [],
                key=lambda a: (
                    str(a.numero_accion)
                    if a.numero_accion
                    else "",
                    a.id
                )
            )

            # ------------------------------------------------
            # Si no tiene acciones, no se genera bloque.
            # ------------------------------------------------

            if not acciones:
                continue

            # =================================================
            # CABECERA DEL SOCIO
            # =================================================

            ws.merge_cells(start_row=fila,start_column=1,end_row=fila,end_column=11)
            celda_socio = ws.cell(row=fila,column=1)
            celda_socio.value = (f"Socio :    "f"{socio.nombres}")
            celda_socio.font = Font(bold=True,size=12)
            celda_socio.fill = PatternFill(fill_type="solid",fgColor=azul_claro)
            celda_socio.alignment = Alignment(horizontal="left")
            fila += 1
            # =================================================
            # CADA ACCIÓN DEL SOCIO
            # =================================================
            for accion in acciones:
                # ------------------------------------------------
                # PRÉSTAMOS DE ESTA ACCIÓN
                # ------------------------------------------------
                prestamos_accion = (prestamos_por_accion.get(accion.id,[]))
                monto_prestamos = sum(
                    (
                        Decimal(
                            p.monto
                        )
                        for p in prestamos_accion
                    ),
                    Decimal("0.00")
                )

                # ------------------------------------------------
                # CUOTA FIJA
                #
                # Si existen varios préstamos, se toma la suma
                # de sus cuotas mínimas.
                # ------------------------------------------------
                cuota_fija = sum(
                    (
                        Decimal(
                            p.cuota_minima
                        )
                        for p in prestamos_accion
                    ),
                    Decimal("0.00")
                )

                # ------------------------------------------------
                # SALDO ACTUAL
                #
                # Se toma el saldo más reciente registrado
                # para esta acción.
                # ------------------------------------------------
                saldo_actual = Decimal("0.00")
                movimientos_accion = []

                for mes in range(1, 13):

                    movimientos_mes = (
                        movimientos_por_accion_mes
                        .get(
                            accion.id,
                            {}
                        )
                        .get(
                            mes,
                            []
                        )
                    )

                    movimientos_accion.extend(
                        movimientos_mes
                    )


                if movimientos_accion:

                    movimientos_accion.sort(
                        key=lambda x: (
                            x.fecha_registro
                            or 0
                        )
                    )

                    ultimo_movimiento = (
                        movimientos_accion[-1]
                    )

                    saldo_actual = Decimal(
                        ultimo_movimiento.saldo_prestamo
                    )


                # =================================================
                # IDENTIFICACIÓN DE ACCIÓN
                # =================================================
                ws.merge_cells(start_row=fila,start_column=1,end_row=fila,end_column=11)
                celda_accion = ws.cell(row=fila,column=1)
                numero_accion = (
                    accion.numero_accion
                    if accion.numero_accion
                    else accion.id
                )

                celda_accion.value = (f"Acción {numero_accion}:")
                celda_accion.font = Font(bold=True)
                celda_accion.fill = PatternFill(fill_type="solid",fgColor=gris)

                fila += 1

                # =================================================
                # CABECERA
                # =================================================
                encabezados = [
                    "Mes",
                    "Aportes",
                    "Amortización",
                    "Intereses",
                    "TOTAL",
                    "Saldo de Préstamos",
                    "Cuota a pagar",
                    "Cuota fija",
                    "Préstamos",
                    "Multa",
                    "Observación"
                ]

                for columna, texto in enumerate(
                    encabezados,
                    start=1
                ):

                    cell = ws.cell(row=fila,column=columna)
                    cell.value = texto
                    cell.font = Font(bold=True,color=negro)
                    cell.fill = PatternFill(fill_type="solid",fgColor=azul)
                    cell.alignment = Alignment(horizontal="center",vertical="center",wrap_text=True)
                    cell.border = border

                fila += 1

                # =================================================
                # INICIO 2026
                # AQUÍ SE USA DIRECTAMENTE accion.valor
                # =================================================

                aporte_inicio = Decimal(accion.valor)

                ws.cell(row=fila,column=1).value = (f"Inicio {anio}")
                ws.cell(row=fila,column=2).value = aporte_inicio
                ws.cell(row=fila,column=3).value = "....."
                ws.cell(row=fila,column=4).value = "....."
                ws.cell(row=fila,column=5).value = "....."
                ws.cell(row=fila,column=6).value = (
                    saldo_actual
                    if saldo_actual > 0
                    else Decimal("0.00")
                )
                ws.cell(row=fila,column=7).value = "....."
                ws.cell(row=fila,column=8).value = "....."
                ws.cell(row=fila,column=9).value = (
                    "Según balance"
                )
                ws.cell(row=fila,column=10).value = ""
                ws.cell(row=fila,column=11).value = ""

                for col in range(1, 12):

                    cell = ws.cell(row=fila,column=col)
                    cell.border = border
                    cell.alignment = Alignment(vertical="center")

                ws.cell(row=fila,column=2).number_format = '#,##0.00'
                ws.cell(row=fila,column=6).number_format = '#,##0.00'

                fila += 1
                # =================================================
                # ACUMULADORES DE LA ACCIÓN
                # =================================================
                total_aportes_accion = Decimal("0.00")
                total_amortizacion_accion = Decimal("0.00")
                total_intereses_accion = Decimal("0.00")
                total_multas_accion = Decimal("0.00")
                total_cuotas_accion = Decimal("0.00")
                ultimo_saldo = Decimal("0.00")
                # =================================================
                # MESES
                # =================================================
                nombres_meses = [
                    "",
                    "Enero",
                    "Febrero",
                    "Marzo",
                    "Abril",
                    "Mayo",
                    "Junio",
                    "Julio",
                    "Agosto",
                    "Setiembre",
                    "Octubre",
                    "Noviembre",
                    "Diciembre"
                ]

                for mes in range(1, 13):

                    lista = (
                        movimientos_por_accion_mes
                        .get(
                            accion.id,
                            {}
                        )
                        .get(
                            mes,
                            []
                        )
                    )

                    aportes = suma_movimientos(lista,"aporte")
                    amortizacion = suma_movimientos(lista,"amortizacion")
                    intereses = suma_movimientos(lista,"interes")
                    cuotas = suma_movimientos(lista,"cuota_pagada")
                    multas = suma_movimientos(lista,"multa")
                    sobres = suma_movimientos(lista,"sobre")

                    # ---------------------------------------------
                    # TOTAL
                    # SOLO:
                    # Aportes + Amortización + Intereses
                    # ---------------------------------------------

                    total_mes = (aportes + amortizacion + intereses)

                    # ---------------------------------------------
                    # SALDO
                    # ---------------------------------------------
                    saldo_mes = ultimo_saldo

                    if lista:

                        ultimo_mov = sorted(
                            lista,
                            key=lambda x: (
                                x.fecha_registro
                                or 0
                            )
                        )[-1]

                        saldo_mes = Decimal(
                            ultimo_mov.saldo_prestamo
                        )

                        ultimo_saldo = saldo_mes

                    # =================================================
                    # ACUMULAR RESUMEN MENSUAL
                    # =================================================

                    resumen_mensual[mes]["aportes"] += aportes
                    resumen_mensual[mes]["amortizacion"] += amortizacion
                    resumen_mensual[mes]["intereses"] += intereses
                    resumen_mensual[mes]["cuota_pagar"] += cuotas
                    resumen_mensual[mes]["multa"] += multas
                    resumen_mensual[mes]["sobre"] += sobres

                    resumen_mensual[mes]["total"] += (
                        aportes +
                        amortizacion +
                        intereses
                    )

                    # ---------------------------------------------
                    # OBSERVACIÓN
                    # ---------------------------------------------
                    observaciones = []

                    for mov in lista:

                        if mov.observacion:

                            observaciones.append(
                                str(
                                    mov.observacion
                                )
                            )


                    observacion = "; ".join(
                        observaciones
                    )

                    # ---------------------------------------------
                    # ESCRIBIR FILA
                    # ---------------------------------------------
                    ws.cell(row=fila,column=1).value = nombres_meses[mes]
                    ws.cell(row=fila,column=2).value = aportes
                    ws.cell(row=fila,column=3).value = amortizacion
                    ws.cell(row=fila,column=4).value = intereses
                    ws.cell(row=fila,column=5).value = total_mes
                    ws.cell(row=fila,column=6).value = saldo_mes
                    ws.cell(row=fila,column=7).value = cuotas
                    ws.cell(row=fila,column=8).value = (
                        cuota_fija
                        if cuota_fija > 0
                        else 0
                    )

                    # ---------------------------------------------
                    # PRÉSTAMO
                    # Se muestra el monto real del Prestamo.
                    # No se suma al TOTAL.
                    # ---------------------------------------------

                    ws.cell(row=fila,column=9).value = (
                        monto_prestamos
                        if mes == 1
                        else 0
                    )
                    ws.cell(row=fila,column=10).value = multas
                    ws.cell(row=fila,column=11).value = observacion

                    # ---------------------------------------------
                    # BORDES
                    # ---------------------------------------------
                    for col in range(1, 12):

                        cell = ws.cell(row=fila,column=col)
                        cell.border = border
                        cell.alignment = Alignment(vertical="center")

                    # ---------------------------------------------
                    # FORMATO NUMÉRICO
                    # ---------------------------------------------
                    for col in range(2, 11):

                        ws.cell(row=fila,column=col).number_format = '#,##0.00'

                    # ---------------------------------------------
                    # ACUMULADORES
                    # ---------------------------------------------
                    total_aportes_accion += aportes
                    total_amortizacion_accion += amortizacion
                    total_intereses_accion += intereses
                    total_multas_accion += multas
                    total_cuotas_accion += cuotas

                    fila += 1

                # =================================================
                # CIERRE DEL AÑO POR ACCIÓN
                # =================================================
                total_accion = (
                    total_aportes_accion +
                    total_amortizacion_accion +
                    total_intereses_accion
                )

                ws.cell(row=fila,column=1).value = (f"Cierre {anio}")
                ws.cell(row=fila,column=2).value = (aporte_inicio + total_aportes_accion)
                ws.cell(row=fila,column=3).value = total_amortizacion_accion
                ws.cell(row=fila,column=4).value = total_intereses_accion
                ws.cell(row=fila,column=5).value = total_accion
                ws.cell(row=fila,column=6).value = ultimo_saldo
                ws.cell(row=fila,column=7).value = total_cuotas_accion
                ws.cell(row=fila,column=8).value = cuota_fija
                ws.cell(row=fila,column=9).value = monto_prestamos
                ws.cell(row=fila,column=10).value = total_multas_accion
                ws.cell(row=fila,column=11).value = ""

                for col in range(1, 12):

                    cell = ws.cell(row=fila,column=col)
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(fill_type="solid",fgColor=celeste)
                    cell.border = border
                    cell.alignment = Alignment(vertical="center")

                for col in range(2, 11):

                    ws.cell(row=fila,column=col).number_format = '#,##0.00'

                # =================================================
                # ACUMULAR TOTALES GENERALES
                # =================================================
                gran_total_aportes += (
                    aporte_inicio +
                    total_aportes_accion
                )

                gran_total_amortizacion += (
                    total_amortizacion_accion
                )

                gran_total_intereses += (
                    total_intereses_accion
                )

                gran_total += (
                    total_accion
                )

                gran_total_prestamos += (
                    monto_prestamos
                )

                gran_total_multas += (
                    total_multas_accion
                )

                fila += 2

        # ====================================================
        # RESUMEN GENERAL MENSUAL
        # ====================================================

        fila += 1

        ws.merge_cells(
            start_row=fila,
            start_column=1,
            end_row=fila,
            end_column=11
        )

        ws.cell(
            row=fila,
            column=1
        ).value = (
            f"RESUMEN GENERAL "
            f"ESTADO DE OPERACIONES {anio}"
        )

        ws.cell(
            row=fila,
            column=1
        ).font = Font(
            bold=True,
            color=blanco,
            size=12
        )

        ws.cell(
            row=fila,
            column=1
        ).fill = PatternFill(
            fill_type="solid",
            fgColor=azul_oscuro
        )

        ws.cell(
            row=fila,
            column=1
        ).alignment = Alignment(
            horizontal="center"
        )

        fila += 1

        # ====================================================
        # CABECERA
        # ====================================================

        resumen_headers = [
            "Mes",
            "Aportes",
            "Amortización",
            "Intereses",
            "TOTAL",
            "Saldo de Préstamos",
            "Cuota a pagar",
            "Cuota fija",
            "Préstamos",
            "Multa",
            "Sobre"
        ]

        for col, texto in enumerate(
            resumen_headers,
            start=1
        ):

            cell = ws.cell(
                row=fila,
                column=col
            )

            cell.value = texto
            cell.font = Font(
                bold=True,
                color=negro
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=azul
            )

            cell.border = border

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

        fila += 1

        # ====================================================
        # ANCHOS
        # ====================================================
        anchos = {
            "A": 22,
            "B": 16,
            "C": 18,
            "D": 16,
            "E": 16,
            "F": 20,
            "G": 16,
            "H": 16,
            "I": 18,
            "J": 14,
            "K": 35
        }

        for columna, ancho in anchos.items():

            ws.column_dimensions[
                columna
            ].width = ancho

        # ====================================================
        # ALINEACIONES
        # ====================================================
        for row in ws.iter_rows():

            for cell in row:

                if cell.column >= 2:

                    if isinstance(
                        cell.value,
                        (int, float, Decimal)
                    ):

                        cell.alignment = Alignment(horizontal="right",vertical="center")

        # ====================================================
        # CONGELAR
        # ====================================================
        ws.freeze_panes = "A3"

        # ====================================================
        # FILTRO
        # ====================================================
        # No se aplica AutoFilter porque el reporte contiene
        # bloques independientes por socio/acción.

        # ====================================================
        # PÁGINA
        # ====================================================
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

        ws.page_setup.orientation = "landscape"

        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5

        # ====================================================
        # GENERAR ARCHIVO
        # ====================================================
        output = BytesIO()

        wb.save(output)

        output.seek(0)

        nombre_archivo = (
            f"BK_Familiar_Estado_Operaciones_"
            f"{anio}.xlsx"
        )

        return send_file(

            output,
            as_attachment=True,
            download_name=nombre_archivo,
            mimetype=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            )

        )

    except Exception as e:

        db.rollback()
        raise e

    finally:

        db.close()

#
# 
# 

